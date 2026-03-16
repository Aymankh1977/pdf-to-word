import io
import re
import os
import streamlit as st
import pdfplumber
from docx import Document
from docx.shared import Pt, Inches

st.set_page_config(page_title="PDF to Word Converter", page_icon="📄", layout="centered")

st.markdown("""
<style>
    .stApp { background-color: #0f0f0f; color: #f0ede8; }
    h1 { color: #c8a96e; }
    .stButton > button {
        background-color: #c8a96e; color: #0f0f0f;
        border: none; border-radius: 8px; width: 100%;
    }
    [data-testid="stDownloadButton"] > button {
        background-color: #1e1e1e; color: #c8a96e;
        border: 1.5px solid #c8a96e; border-radius: 8px; width: 100%;
    }
    #MainMenu, footer { visibility: hidden; }
</style>
""", unsafe_allow_html=True)

st.title("PDF → Word Converter")
st.write("Upload one or more PDFs and download clean Word documents.")


def clean_text(text):
    if not text:
        return ""
    text = text.replace('\xa0', ' ')
    text = re.sub(r'[ \t]{2,}', ' ', text)
    return text.strip()


def looks_like_heading(text, font_size=None, is_bold=False):
    s = text.strip()
    if not s:
        return 0
    if s.isupper() and 3 <= len(s) <= 80:
        return 1
    if is_bold and len(s) <= 80 and not s.endswith('.'):
        return 2
    if font_size:
        if font_size >= 18:
            return 1
        if font_size >= 14:
            return 2
        if font_size >= 12 and is_bold:
            return 3
    return 0


def extract_page_content(page):
    content = []
    tables = page.extract_tables()
    try:
        table_bboxes = [t.bbox for t in page.find_tables()]
    except Exception:
        table_bboxes = []

    words = page.extract_words(extra_attrs=["size", "fontname"])
    if not words:
        raw = page.extract_text()
        if raw:
            for line in raw.splitlines():
                line = clean_text(line)
                if line:
                    content.append({"type": "paragraph", "text": line, "font_size": None, "bold": False})
        for tbl in tables:
            if tbl:
                content.append({"type": "table", "data": tbl})
        return content

    lines = {}
    for w in words:
        key = round(w['top'], 1)
        lines.setdefault(key, []).append(w)

    def in_table(y):
        for bbox in table_bboxes:
            if bbox[1] <= y <= bbox[3]:
                return True
        return False

    prev_bottom = None
    for top, line_words in sorted(lines.items()):
        if in_table(top):
            continue
        line_words.sort(key=lambda w: w['x0'])
        line_text = clean_text(' '.join(w['text'] for w in line_words))
        if not line_text:
            continue
        sizes = [w.get('size', 0) for w in line_words if w.get('size')]
        font_size = max(sizes) if sizes else None
        fontnames = [w.get('fontname', '') for w in line_words]
        is_bold = any('Bold' in fn or 'bold' in fn for fn in fontnames)
        gap = (top - prev_bottom) if prev_bottom is not None else 0
        if gap > 12:
            content.append({"type": "blank"})
        content.append({"type": "paragraph", "text": line_text, "font_size": font_size, "bold": is_bold})
        prev_bottom = top + line_words[0].get('height', 10)

    for tbl in tables:
        if tbl:
            content.append({"type": "table", "data": tbl})

    return content


def add_table_to_doc(doc, table_data):
    rows = [r for r in table_data if any(c for c in r)]
    if not rows:
        return
    num_cols = max(len(r) for r in rows)
    tbl = doc.add_table(rows=len(rows), cols=num_cols)
    tbl.style = 'Table Grid'
    for r_idx, row in enumerate(rows):
        for c_idx, cell_text in enumerate(row):
            if c_idx < num_cols:
                cell = tbl.cell(r_idx, c_idx)
                cell.text = clean_text(cell_text or "")
                if r_idx == 0:
                    for run in cell.paragraphs[0].runs:
                        run.bold = True


def convert_pdf_to_docx(pdf_bytes):
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1.2)
        section.right_margin = Inches(1.2)
    doc.styles['Normal'].font.name = 'Calibri'
    doc.styles['Normal'].font.size = Pt(11)

    with pdfplumber.open(io.BytesIO(pdf_bytes)) as pdf:
        total = len(pdf.pages)
        for i, page in enumerate(pdf.pages):
            for item in extract_page_content(page):
                if item["type"] == "blank":
                    continue
                elif item["type"] == "paragraph":
                    text = item["text"]
                    level = looks_like_heading(text, item.get("font_size"), item.get("bold", False))
                    if level in (1, 2, 3):
                        doc.add_heading(text, level=level)
                    else:
                        p = doc.add_paragraph(text)
                        if item.get("bold"):
                            for run in p.runs:
                                run.bold = True
                elif item["type"] == "table":
                    add_table_to_doc(doc, item["data"])
            if i < total - 1:
                doc.add_page_break()

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


uploaded_files = st.file_uploader(
    "Upload PDF files",
    type=["pdf"],
    accept_multiple_files=True
)

if uploaded_files:
    st.write(f"{len(uploaded_files)} file(s) ready to convert.")
    if st.button("Convert to Word"):
        for f in uploaded_files:
            with st.spinner(f"Converting {f.name}..."):
                try:
                    docx_bytes = convert_pdf_to_docx(f.read())
                    out_name = os.path.splitext(f.name)[0] + ".docx"
                    st.success(f"✅ {f.name} converted!")
                    st.download_button(
                        label=f"⬇ Download {out_name}",
                        data=docx_bytes,
                        file_name=out_name,
                        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                        key=out_name
                    )
                except Exception as e:
                    st.error(f"❌ Failed to convert {f.name}: {e}")
